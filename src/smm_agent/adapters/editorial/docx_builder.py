"""DOCX preview generated from canonical editorial artifacts."""

from io import BytesIO

from docx import Document as DocxDocument
from docx.document import Document
from docx.shared import Inches

from smm_agent.contracts.editorial import EditorialBundle


class PythonDocxBuilder:
    def build(self, bundle: EditorialBundle, assets: dict[str, bytes]) -> bytes:
        document = DocxDocument()
        document.add_heading(bundle.dzen.title, level=0)
        document.add_picture(BytesIO(assets[bundle.cover_path]), width=Inches(6.0))

        cursor = 0
        for visual in bundle.visuals:
            anchor_end = bundle.main_text.index(visual.anchor_text, cursor) + len(
                visual.anchor_text
            )
            self._add_text(document, bundle.main_text[cursor:anchor_end])
            document.add_picture(BytesIO(assets[visual.asset_path]), width=Inches(5.5))
            caption = document.add_paragraph(visual.caption)
            caption.style = document.styles["Caption"]
            cursor = anchor_end
        self._add_text(document, bundle.main_text[cursor:])

        output = BytesIO()
        document.save(output)
        return output.getvalue()

    @staticmethod
    def _add_text(document: Document, text: str) -> None:
        for paragraph in (part.strip() for part in text.split("\n\n")):
            if paragraph:
                document.add_paragraph(paragraph)
